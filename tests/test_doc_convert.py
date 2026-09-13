from pathlib import Path

import pytest
from docx import Document

from bot.services.doc_convert import ConversionError, convert_document


def _make_docx(path: Path) -> Path:
    document = Document()
    document.add_paragraph("Hello world")
    document.add_paragraph("")
    document.add_paragraph("Wörld ünïcode 你好")
    document.save(str(path))
    return path


def test_docx_to_txt_preserves_paragraphs(settings, tmp_path: Path) -> None:
    src = _make_docx(tmp_path / "report.docx")
    out = tmp_path / "report.txt"
    count = convert_document(src, out, settings=settings)
    text = out.read_text(encoding="utf-8")
    assert count >= 3
    assert "Hello world" in text
    assert "Wörld ünïcode 你好" in text
    assert "\n\n" in text  # blank paragraph preserved


def test_unsupported_extension_raises(settings, tmp_path: Path) -> None:
    src = tmp_path / "file.pdf"
    src.write_bytes(b"%PDF-1.4")
    with pytest.raises(ConversionError):
        convert_document(src, tmp_path / "out.txt", settings=settings)


def test_corrupt_docx_raises(settings, tmp_path: Path) -> None:
    src = tmp_path / "broken.docx"
    src.write_bytes(b"not a real docx")
    with pytest.raises(ConversionError):
        convert_document(src, tmp_path / "out.txt", settings=settings)
